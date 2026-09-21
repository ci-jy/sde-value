"""Render workbook tabs, the dashboard PDF, a memo and the batch summary to PNG for
the README (headless LibreOffice + pdftoppm + ImageMagick).

    python3 scripts/screenshots.py out/hvac_service.xlsx docs/ [out/hvac_service-memo.docx out/summary.xlsx]
"""

import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from sde_value.check import soffice_binary
from sde_value.export import export_dashboard

PAGES = {"Add-back Register": "addback-register.png", "Normalised Earnings": "normalised-earnings.png",
         "Multiple Valuation": "multiple-valuation.png", "DCF": "dcf.png",
         "Financing": "financing.png", "Sensitivity": "sensitivity.png"}


def to_pdf(src: Path, tmp: Path) -> Path:
    subprocess.run([soffice_binary(), "--headless", f"-env:UserInstallation={(tmp / 'p').as_uri()}",
                    "--convert-to", "pdf", "--outdir", str(tmp), str(src)],
                   check=True, capture_output=True)
    return tmp / f"{src.stem}.pdf"


def png(pdf: Path, page: int, dest: Path, tmp: Path, trim: bool = True) -> None:
    stem = tmp / f"{pdf.stem}-{page}"
    subprocess.run(["pdftoppm", "-r", "200" if trim else "110", "-png", "-singlefile",
                    "-f", str(page), "-l", str(page), str(pdf), str(stem)], check=True)
    if trim:
        subprocess.run(["convert", f"{stem}.png", "-trim", "-bordercolor", "white",
                        "-border", "20", "-strip", str(dest)], check=True)
    else:
        subprocess.run(["convert", f"{stem}.png", "-bordercolor", "#cccccc", "-border", "1",
                        "-strip", str(dest)], check=True)
    print(dest)


def main(src: str, out_dir: str, memo: str | None = None, summary: str | None = None) -> None:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        wb = load_workbook(src)
        for ws in wb.worksheets:
            ws.page_setup.orientation = "landscape"
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
        wb.save(tmp / "print.xlsx")
        pdf = to_pdf(tmp / "print.xlsx", tmp)
        for page, name in PAGES.items():
            png(pdf, wb.sheetnames.index(page) + 1, out / name, tmp)

        dash = export_dashboard(src, out / "dashboard.pdf")
        print(dash)
        png(dash, 1, out / "dashboard.png", tmp)
        if memo:
            png(to_pdf(Path(memo), tmp), 1, out / "memo.png", tmp, trim=False)
        if summary:
            png(to_pdf(Path(summary), tmp), 1, out / "summary.png", tmp)


if __name__ == "__main__":
    main(*sys.argv[1:5])
