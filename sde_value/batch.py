"""Value every deal file in a folder: a workbook and memo for each, and one
summary workbook that lines the deals up side by side."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .deal import Deal, DealError, load_deal
from .engine import VERDICTS, value_deal
from .memo import write_memo
from .workbook import BOLD, HEAD_FILL, HEAD_FONT, MONEY, MULT, NOTE, PCT, TITLE, TOP, \
    build_workbook

# (key, header, width, number format). Keys ending in "=" are formulas in the
# summary sheet, written by _formula(); the rest are engine values.
COLUMNS = [
    ("name", "Deal", 34, None),
    ("file", "Deal file", 20, None),
    ("state", "State", 7, None),
    ("revenue", "Revenue", 13, MONEY),
    ("asking_multiple", "Asking multiple", 10, MULT),
    ("seller_sde", "Seller's SDE", 13, MONEY),
    ("rejected_total", "Add-backs rejected", 13, MONEY),
    ("normalised_ebitda", "Normalised EBITDA", 13, MONEY),
    ("asking_value=", "Asking value", 14, MONEY),
    ("repriced_value=", "Repriced value", 14, MONEY),
    ("dcf_value", "DCF value", 13, MONEY),
    ("fin_max_price", "Largest financeable", 14, MONEY),
    ("supported_value=", "Supported value", 14, MONEY),
    ("supported_vs_asking_pct=", "Supported vs asking", 11, PCT),
    ("fin_asking_dscr", "DSCR at asking", 9, '0.00"x"'),
    ("fin_repriced_dscr", "DSCR at repriced", 9, '0.00"x"'),
    ("financing_verdict", "Financing verdict", 40, None),
    ("workbook", "Workbook", 22, None),
    ("memo", "Memo", 26, None),
]
COL = {key.rstrip("="): get_column_letter(i) for i, (key, *_) in enumerate(COLUMNS, 1)}


@dataclass
class DealResult:
    deal: Deal
    outputs: dict
    workbook: Path
    memo: Path
    check: object = None    # CheckResult when run with check=True


@dataclass
class BatchResult:
    results: list[DealResult] = field(default_factory=list)
    errors: list[tuple[Path, str]] = field(default_factory=list)
    summary: Path | None = None


def _formula(key: str, r: int) -> str:
    c = COL
    return {
        "asking_value": f"={c['seller_sde']}{r}*{c['asking_multiple']}{r}",
        "repriced_value": f"={c['normalised_ebitda']}{r}*{c['asking_multiple']}{r}",
        "supported_value": f"=MAX(0,MIN({c['repriced_value']}{r},{c['fin_max_price']}{r}))",
        "supported_vs_asking_pct": (f"=IF({c['asking_value']}{r}=0,0,"
                                    f"{c['supported_value']}{r}/{c['asking_value']}{r}-1)"),
    }[key]


def write_summary(results: list[DealResult], out_path: str | Path) -> Path:
    out_path = Path(out_path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Deal comparison"
    ws["A1"].font = TITLE
    ws["A2"] = ("Asking value is the broker's multiple on the seller's SDE; repriced value the "
                "same multiple on normalised EBITDA. Supported value is the lower of the "
                "repriced value and the largest price the cash flow can finance.")
    ws["A2"].font = NOTE
    head = 4
    for i, (_, label, width, _) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=head, column=i, value=label)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[head].height = 44
    for r, res in enumerate(results, head + 1):
        o, d = res.outputs, res.deal
        values = {"name": d.name, "file": d.path.name if d.path else "", "state": d.state,
                  "revenue": d.revenue, "asking_multiple": d.asking_multiple,
                  "workbook": res.workbook.name, "memo": res.memo.name}
        for i, (key, _, _, fmt) in enumerate(COLUMNS, 1):
            if key.endswith("="):
                value = _formula(key[:-1], r)
            else:
                value = values[key] if key in values else o[key]
            cell = ws.cell(row=r, column=i, value=value)
            if fmt:
                cell.number_format = fmt
            if key in ("workbook", "memo"):
                # Relative links work wherever the out/ folder is moved.
                cell.hyperlink = value
                cell.font = Font(color="0563C1", underline="single")
    first, last = head + 1, head + max(1, len(results))
    avg = last + 1
    ws.cell(row=avg, column=1, value=f"Average of {len(results)} deals").font = BOLD
    col = COL["supported_vs_asking_pct"]
    cell = ws[f"{col}{avg}"]
    cell.value = f"=AVERAGE({col}{first}:{col}{last})" if results else 0
    cell.number_format, cell.font, cell.border = PCT, BOLD, TOP
    v = COL["financing_verdict"]
    ws.conditional_formatting.add(f"{v}{first}:{v}{last}", FormulaRule(
        formula=[f'{v}{first}="{VERDICTS["neither"]}"'],
        fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(f"{v}{first}:{v}{last}", FormulaRule(
        formula=[f'{v}{first}="{VERDICTS["asking"]}"'],
        fill=PatternFill("solid", fgColor="E2EFDA")))
    ws.freeze_panes = f"B{head + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def run_batch(deals_dir: str | Path, out_dir: str | Path, check: bool = False) -> BatchResult:
    """Build a workbook and memo per deal file, then the summary workbook.

    A deal file that fails to load is reported and skipped; the rest still run.
    """
    deals_dir, out_dir = Path(deals_dir), Path(out_dir)
    paths = sorted(p for p in deals_dir.iterdir() if p.suffix in (".yaml", ".yml"))
    if not paths:
        raise DealError(f"no .yaml deal files in {deals_dir}")
    batch = BatchResult()
    for path in paths:
        try:
            deal = load_deal(path)
            outputs = value_deal(deal)
        except (DealError, KeyError, ValueError) as exc:
            batch.errors.append((path, str(exc)))
            continue
        xlsx = build_workbook(deal, out_dir / f"{deal.slug}.xlsx")
        memo = write_memo(deal, out_dir / f"{deal.slug}-memo.docx")
        res = DealResult(deal, outputs, xlsx, memo)
        if check:
            from .check import check_workbook
            res.check = check_workbook(xlsx)
        batch.results.append(res)
    batch.summary = write_summary(batch.results, out_dir / "summary.xlsx")
    return batch
