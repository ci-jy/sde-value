"""Print the workbook's dashboard tab to PDF with headless LibreOffice."""

from __future__ import annotations

import re
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from .check import RecalcError, soffice_binary
from .workbook import S_DASH


def export_dashboard(workbook: str | Path, pdf_path: str | Path) -> Path:
    """Write a one-page PDF of the Dashboard tab.

    LibreOffice prints every visible sheet, so a copy of the workbook is saved
    with the other tabs hidden; their formulas still feed the dashboard.
    """
    pdf_path = Path(pdf_path)
    with tempfile.TemporaryDirectory(prefix="sde-pdf-") as tmp:
        tmp = Path(tmp)
        wb = load_workbook(workbook)
        if S_DASH not in wb.sheetnames:
            raise RecalcError(f"{workbook} has no {S_DASH} tab; rebuild it with `build`")
        wb.active = wb.sheetnames.index(S_DASH)
        for ws in wb.worksheets:
            if ws.title != S_DASH:
                ws.sheet_state = "hidden"
        copy = tmp / "dashboard.xlsx"
        wb.save(copy)
        cmd = [soffice_binary(), "--headless", "--norestore", "--nolockcheck",
               f"-env:UserInstallation={(tmp / 'lo-profile').as_uri()}",
               "--convert-to", "pdf", "--outdir", str(tmp), str(copy)]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=180)
        result = tmp / "dashboard.pdf"
        if proc.returncode != 0 or not result.exists():
            raise RecalcError(f"LibreOffice failed: {proc.stderr.strip() or proc.stdout.strip()}")
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        pdf_path.write_bytes(_blank_timestamps(result.read_bytes()))
    return pdf_path


def _blank_timestamps(pdf: bytes) -> bytes:
    """Blank the creation/modification times LibreOffice stamps into the PDF.

    The PDF is committed to docs/, and a re-export shouldn't differ just by
    the time it ran. Each stamp is overwritten with spaces of the same length,
    so the xref offsets still hold.
    """
    blank = lambda m: m.group(1) + b" " * len(m.group(2)) + m.group(3)  # noqa: E731
    pdf = re.sub(rb"(/(?:CreationDate|ModDate)\()(D:[^)]*)(\))", blank, pdf)
    return re.sub(rb"(<xmp:\w*Date>)([^<]+)(</)", blank, pdf)
