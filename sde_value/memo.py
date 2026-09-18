"""A short Word memo per deal: asking vs supported value, the add-backs that
were rejected and why, and whether the price can be financed.

Every number comes from the engine, so the memo says what the workbook says
when the deal file's toggles are left as they are.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from .addbacks import classify_addback, rejection_reason
from .deal import Deal
from .engine import PRICE_CASES, value_deal


@dataclass
class Rejected:
    item: str
    amount: float
    classification: str
    reason: str


@dataclass
class MemoFacts:
    deal: Deal
    outputs: dict
    rejected: list[Rejected]
    overridden: list[str]     # add-backs the buyer accepted against the rule


def memo_facts(deal: Deal) -> MemoFacts:
    rejected, overridden = [], []
    for ab in deal.addbacks:
        cls, recommended = classify_addback(ab)
        reason = rejection_reason(ab)
        if reason:
            rejected.append(Rejected(ab.item, ab.amount, cls, reason))
        elif not recommended:
            overridden.append(ab.item)
    return MemoFacts(deal, value_deal(deal), rejected, overridden)


def _money(x: float) -> str:
    return f"-${-x:,.0f}" if x < 0 else f"${x:,.0f}"


def _pct(x: float) -> str:
    return f"{x:+.1%}".replace("+0.0%", "0.0%")


def bottom_line(f: MemoFacts) -> str:
    o, d = f.outputs, f.deal
    n = len(f.rejected)
    text = (f"The broker asks {_money(o['asking_value'])}, {d.asking_multiple:.1f}x a seller's "
            f"SDE of {_money(o['seller_sde'])}. ")
    if n:
        text += (f"After rejecting {n} add-back{'s' if n != 1 else ''} worth "
                 f"{_money(o['rejected_total'])} and ")
    else:
        text += "Every add-back holds up, but after "
    text += (f"paying a manager {_money(o['replacement_cost'])} a year, the same multiple gives "
             f"{_money(o['repriced_value'])}. ")
    if o["supported_basis"] == "repriced value":
        text += "The cash flow can finance that, so it is the supported value"
    else:
        text += (f"The cash flow can only carry a loan on {_money(o['fin_max_price'])}, "
                 "so that is the supported value")
    gap = o["supported_vs_asking_pct"]
    text += (f": {_money(o['supported_value'])}, {abs(gap):.1%} "
             f"{'below' if gap < 0 else 'above'} the asking price.")
    return text


def financing_text(f: MemoFacts) -> str:
    o, loan = f.outputs, f.deal.loan
    bank = f"a bank loan over {loan.senior_years:g} years at {loan.senior_rate:.2%}"
    if loan.seller_note:
        terms = (f"{loan.down_payment:.0%} down, {bank} and a {loan.seller_note:.0%} seller "
                 f"note over {loan.note_years:g} years at {loan.note_rate:.2%}")
    else:
        terms = f"{loan.down_payment:.0%} down and {bank}"
    return (f"With {terms}, the buyer's cash flow of "
            f"{_money(o['fin_cfads'])} (normalised EBITDA less maintenance capex) covers the "
            f"debt service {o['fin_asking_dscr']:.2f}x at the asking price and "
            f"{o['fin_repriced_dscr']:.2f}x at the repriced value. Lenders look for "
            f"{loan.min_dscr:.2f}x, which caps the price at {_money(o['fin_max_price'])}.")


def _table(doc, header: list[str], rows: list[list[str]], bold_last: bool = False,
           right: tuple[int, ...] = (), widths: tuple[float, ...] = ()):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.autofit = not widths
    for cell, text in zip(t.rows[0].cells, header):
        cell.text = text
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    for r_i, row in enumerate(t.rows):
        for c_i, cell in enumerate(row.cells):
            for par in cell.paragraphs:
                if c_i in right:
                    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if bold_last and r_i == len(t.rows) - 1:
                    for run in par.runs:
                        run.bold = True
            if widths:
                cell.width = Inches(widths[c_i])
    for c_i, w in enumerate(widths):
        t.columns[c_i].width = Inches(w)
    # Word runs a paragraph straight into the table above it; leave a gap.
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def write_memo(deal: Deal, out_path: str | Path) -> Path:
    f = memo_facts(deal)
    o = f.outputs
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(60)
        section.top_margin = section.bottom_margin = Pt(54)

    doc.add_heading(f"Buyer memo: {deal.name}", level=1)
    sub = doc.add_paragraph(deal.description)
    sub.runs[0].italic = True

    doc.add_heading("Bottom line", level=2)
    doc.add_paragraph(bottom_line(f))

    doc.add_heading("Asking vs supported value", level=2)
    sde = o["seller_sde"]
    rows = []
    for case, label in PRICE_CASES.items():
        v = o[f"{case}_value"]
        rows.append([label, _money(v), _pct(v / o["asking_value"] - 1 if o["asking_value"] else 0),
                     f"{v / sde:.1f}x" if sde else ""])
    rows.append(["Largest financeable price", _money(o["fin_max_price"]),
                 _pct(o["fin_max_vs_asking_pct"]), f"{o['fin_max_seller_multiple']:.1f}x"])
    rows.append(["Supported value", _money(o["supported_value"]),
                 _pct(o["supported_vs_asking_pct"]),
                 f"{o['supported_value'] / sde:.1f}x" if sde else ""])
    _table(doc, ["", "Value", "vs asking", "x seller's SDE"], rows, bold_last=True,
           right=(1, 2, 3), widths=(2.5, 1.4, 1.4, 1.4))

    doc.add_heading("Earnings bridge", level=2)
    _table(doc, ["", "Amount"], [
        ["Seller's SDE (every add-back claimed)", _money(o["seller_sde"])],
        ["Less rejected add-backs", _money(-o["rejected_total"])],
        ["Normalised SDE", _money(o["normalised_sde"])],
        ["Less a manager at market pay and benefits", _money(-o["replacement_cost"])],
        ["Normalised EBITDA", _money(o["normalised_ebitda"])],
    ], bold_last=True, right=(1,), widths=(4.3, 2.4))
    doc.add_paragraph(
        f"Normalised EBITDA is {o['normalised_ebitda_margin']:.1%} of revenue against "
        f"{o['base_industry_ebitda_margin']:.1%} for the industry in IRS SOI figures; the "
        f"seller's recast claims {o['seller_sde_margin']:.1%}.")

    doc.add_heading("Rejected add-backs", level=2)
    if f.rejected:
        rows = [[r.item, _money(r.amount), r.classification, r.reason] for r in f.rejected]
        rows.append(["Total", _money(o["rejected_total"]), "", ""])
        _table(doc, ["Item", "Amount", "Classified as", "Why it was rejected"], rows,
               bold_last=True, right=(1,), widths=(1.7, 0.85, 1.25, 2.9))
        doc.add_paragraph(
            f"At {deal.asking_multiple:.1f}x these account for "
            f"{_money(o['value_from_rejected'])} of the asking price.")
    else:
        doc.add_paragraph("None: every add-back the seller claims is accepted.")
    if f.overridden:
        doc.add_paragraph("Accepted although the rules would reject them: "
                          + "; ".join(f.overridden) + ".")

    doc.add_heading("Financing", level=2)
    par = doc.add_paragraph()
    par.add_run(f"Verdict: {o['financing_verdict']}. ").bold = True
    par.add_run(financing_text(f))

    note = doc.add_paragraph(
        "A first-pass screen built from the seller's P&L, public BLS wage data and IRS SOI "
        "industry ratios. It is not an appraisal, tax structuring or legal advice. The "
        "workbook has every figure as a formula, with the add-back toggles to change.")
    for run in note.runs:
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
