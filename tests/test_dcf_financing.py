"""DCF, acquisition loan coverage and the two-way sensitivity table.

The engine is checked by hand on small numbers, then the DCF, Financing and
Sensitivity tabs are recalculated in LibreOffice and matched to the engine.
"""

import shutil

import pytest
from openpyxl import load_workbook

from sde_value.check import RecalcError, check_workbook, compare, named_cells, recalculate
from sde_value.cli import main
from sde_value.deal import DCFTerms, DealError, LoanTerms, load_deal, parse_deal
from sde_value.engine import evaluate, payment_factor, value_deal
from sde_value.workbook import build_workbook

from test_engine import simple_inputs

NEW_TABS = ["DCF", "Financing", "Sensitivity"]

needs_lo = pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")),
                              reason="LibreOffice is required for the recalculation check")


def cell(wb, name):
    sheet, ref = next(iter(wb.defined_names[name].destinations))
    return wb[sheet][ref.replace("$", "")]


# --- engine ----------------------------------------------------------------

def test_payment_factor_is_a_monthly_annuity():
    # $100k over 10 years at 12%: textbook monthly payment of $1,434.71
    assert payment_factor(0.12, 10) * 100_000 / 12 == pytest.approx(1434.71, abs=0.01)
    assert payment_factor(0.0, 5) == pytest.approx(0.2)


def test_dcf_by_hand():
    dcf = DCFTerms(years=2, growth=0.10, capex_pct=0.02, nwc_pct=0.10, tax_rate=0.25,
                   discount_rate=0.20, terminal_growth=0.0)
    o = evaluate(simple_inputs(dcf=dcf))
    margin = 83_000 / 1_000_000          # normalised EBITDA margin from simple_inputs
    rev1, rev2 = 1_100_000, 1_210_000
    fcf1 = rev1 * margin - rev1 * 0.02 - 100_000 * 0.10 - (rev1 * margin - rev1 * 0.02) * 0.25
    fcf2 = rev2 * margin - rev2 * 0.02 - 110_000 * 0.10 - (rev2 * margin - rev2 * 0.02) * 0.25
    assert o["dcf_y1_fcf"] == pytest.approx(fcf1)
    assert o["dcf_y2_fcf"] == pytest.approx(fcf2)
    terminal = fcf2 / 0.20
    value = fcf1 / 1.2 + fcf2 / 1.44 + terminal / 1.44
    assert o["dcf_terminal_value"] == pytest.approx(terminal)
    assert o["dcf_value"] == pytest.approx(value)
    assert o["dcf_ebitda_multiple"] == pytest.approx(value / 83_000)


def test_no_tax_on_a_loss():
    dcf = DCFTerms(years=1, capex_pct=0.5)
    o = evaluate(simple_inputs(dcf=dcf))
    assert o["dcf_y1_tax"] == 0.0
    assert o["dcf_y1_fcf"] < 0


def test_debt_service_and_coverage():
    loan = LoanTerms(down_payment=0.10, seller_note=0.10, senior_rate=0.10, senior_years=10,
                     note_rate=0.06, note_years=5, min_dscr=1.25)
    o = evaluate(simple_inputs(loan=loan, dcf=DCFTerms(capex_pct=0.01)))
    price = o["asking_value"]
    service = price * 0.8 * payment_factor(0.10, 10) + price * 0.1 * payment_factor(0.06, 5)
    assert o["fin_asking_senior"] == pytest.approx(price * 0.8)
    assert o["fin_asking_debt_service"] == pytest.approx(service)
    assert o["fin_cfads"] == pytest.approx(83_000 - 10_000)
    assert o["fin_asking_dscr"] == pytest.approx(73_000 / service)


def test_largest_price_sits_exactly_on_the_minimum_dscr():
    o = evaluate(simple_inputs(dcf=DCFTerms(capex_pct=0.01)))
    service = o["fin_max_price"] * o["fin_service_factor"]
    assert o["fin_cfads"] / service == pytest.approx(1.25)
    assert o["fin_max_seller_multiple"] == pytest.approx(o["fin_max_price"] / o["seller_sde"])


def test_negative_cash_flow_finances_nothing():
    o = evaluate(simple_inputs(industry_wage=400_000, national_wage=400_000, state_wage=400_000))
    assert o["fin_cfads"] < 0
    assert o["fin_max_price"] == 0
    assert o["fin_asking_financeable"] == "no"


def test_sensitivity_corners_are_the_buyer_and_the_broker():
    o = evaluate(simple_inputs(acceptance=[0.0, 0.5, 1.0]))
    # multiples are [3.0, 4.0]; column 2 is the asking multiple
    assert o["sens_1_2_value"] == pytest.approx(o["repriced_value"])
    assert o["sens_3_2_value"] == pytest.approx((o["normalised_ebitda"] + 10_000) * 4)
    assert o["sens_2_1_value"] == pytest.approx((o["normalised_ebitda"] + 5_000) * 3)
    # paying more for the same cash flow can only lower coverage
    assert o["sens_1_1_dscr"] > o["sens_1_2_dscr"] > o["sens_3_2_dscr"]


def test_deal_terms_are_validated():
    base = {
        "name": "x", "industry": {"soi": "plumbing_hvac", "naics": "238200"}, "state": "TX",
        "asking": {"multiple": 4}, "pnl": {"revenue": 100_000, "rent": 10_000},
    }
    assert parse_deal(base).loan.min_dscr == 1.25
    with pytest.raises(DealError, match="terminal_growth"):
        parse_deal({**base, "dcf": {"discount_rate": 0.02, "terminal_growth": 0.03}})
    with pytest.raises(DealError, match="unknown fields"):
        parse_deal({**base, "financing": {"balloon": 0.5}})
    with pytest.raises(DealError, match="bank loan"):
        parse_deal({**base, "financing": {"down_payment": 0.6, "seller_note": 0.4}})
    with pytest.raises(DealError, match="acceptance"):
        parse_deal({**base, "sensitivity": {"acceptance": [0, 1.5]}})


def test_capex_defaults_to_the_sellers_depreciation(hvac):
    from sde_value.engine import inputs_from_deal
    m, _ = inputs_from_deal(hvac)
    assert m.dcf.capex_pct == pytest.approx(42_700 / 2_400_000, abs=0.00005)


def test_hvac_results(hvac):
    """The numbers quoted in the README."""
    o = value_deal(hvac)
    assert o["dcf_value"] == pytest.approx(486_172, abs=1)
    assert o["dcf_ebitda_multiple"] == pytest.approx(3.12, abs=0.005)
    assert o["fin_max_price"] == pytest.approx(593_621, abs=1)
    assert o["fin_max_seller_multiple"] == pytest.approx(1.53, abs=0.005)
    assert o["fin_asking_dscr"] == pytest.approx(0.48, abs=0.005)
    assert (o["fin_asking_financeable"], o["fin_repriced_financeable"],
            o["fin_dcf_financeable"]) == ("no", "no", "yes")
    # 12 of the 35 price/multiple cells clear 1.25x, and none at the 4.0x asking multiple
    cells = [o[f"sens_{k}_{j}_dscr"] for k in range(1, 6) for j in range(1, 8)]
    assert sum(d >= 1.25 for d in cells) == 12
    assert max(o[f"sens_{k}_5_dscr"] for k in range(1, 6)) < 1.25


def test_no_sample_deal_can_borrow_its_asking_price(deal_path):
    o = value_deal(load_deal(deal_path))
    assert o["fin_asking_dscr"] < 1
    assert o["fin_max_price"] < o["repriced_value"]


# --- workbook and LibreOffice -----------------------------------------------

def test_new_tabs_are_live_formulas(tmp_path, hvac):
    path = build_workbook(hvac, tmp_path / "hvac.xlsx")
    wb = load_workbook(path)
    assert wb.sheetnames[4:7] == NEW_TABS
    outputs = [n for n, dn in wb.defined_names.items()
               if n.startswith("o_") and next(iter(dn.destinations))[0] in NEW_TABS]
    assert len(outputs) == 48 + 36 + 75
    for name in outputs:
        value = cell(wb, name).value
        assert isinstance(value, str) and value.startswith("="), name
    # sensitivity cells are formulas over the row/column headers, not pasted numbers
    assert cell(wb, "o_sens_1_1_value").value.startswith("=$B")
    assert "PMT(" in cell(wb, "o_fin_senior_factor").value


@needs_lo
def test_new_tabs_match_engine_after_recalculation(tmp_path, deal_path):
    deal = load_deal(deal_path)
    path = build_workbook(deal, tmp_path / f"{deal.slug}.xlsx")
    result = check_workbook(path, tabs=NEW_TABS)
    assert result.ok, (result.mismatches, result.missing)
    if deal.slug == "hvac_service":
        assert result.compared == 48 + 36 + 75
    assert result.workbook["o_fin_asking_financeable"] == "no"


@needs_lo
def test_edits_in_excel_flow_through(tmp_path, hvac):
    """Change loan terms, the discount rate and a toggle; the tabs still agree."""
    path = build_workbook(hvac, tmp_path / "hvac.xlsx")
    before, _ = named_cells(recalculate(path, tmp_path / "before"))

    wb = load_workbook(path)
    cell(wb, "i_fin_senior_rate").value = 0.08
    cell(wb, "i_fin_seller_note").value = 0.0
    cell(wb, "i_dcf_discount_rate").value = 0.15
    cell(wb, "i_sens_acc2").value = 0.4
    idx = next(i for i, ab in enumerate(hvac.addbacks, 1) if "Las Vegas" in ab.item)
    cell(wb, f"i_ab{idx}_accept").value = 1
    wb.save(path)

    after, sheets = named_cells(recalculate(path, tmp_path / "after"))
    result = compare(after, sheets, tabs=NEW_TABS)
    assert result.ok, result.mismatches
    assert after["o_dcf_value"] > before["o_dcf_value"]
    assert after["o_fin_max_price"] > before["o_fin_max_price"]
    # accepting the trip lowers the rejected pool by 5,200
    assert after["o_sens_5_ebitda"] - after["o_sens_1_ebitda"] == pytest.approx(92_700 - 5_200)


@needs_lo
def test_checker_catches_broken_new_formulas(tmp_path, hvac):
    path = build_workbook(hvac, tmp_path / "hvac.xlsx")
    wb = load_workbook(path)
    cell(wb, "o_dcf_y3_tax").value = "=0"                 # someone drops the tax line
    cell(wb, "o_sens_2_4_value").value = "=500000"        # a pasted number in the grid
    wb.save(path)
    result = check_workbook(path, tabs=NEW_TABS)
    names = {m.name for m in result.mismatches}
    assert {"o_dcf_y3_tax", "o_dcf_y3_fcf", "o_dcf_value", "o_sens_2_4_value",
            "o_sens_2_4_dscr", "o_fin_dcf_dscr"} <= names
    # the earlier tabs are untouched, so a check limited to them still passes
    assert check_workbook(path, tabs=["Normalised Earnings", "Multiple Valuation"]).ok


@needs_lo
def test_unknown_tab_is_an_error(tmp_path, hvac, capsys):
    path = build_workbook(hvac, tmp_path / "hvac.xlsx")
    with pytest.raises(RecalcError, match="Cashflow"):
        check_workbook(path, tabs=["Cashflow"])
    assert main(["check", str(path), "--tabs", "DCF,Cashflow"]) == 2
    assert main(["check", str(path), "--tabs", "DCF"]) == 0
    assert "48 output cells on DCF" in capsys.readouterr().out
