"""Rejection reasons and the Word memo: same numbers as the engine, in sentences."""

import yaml
import pytest
from docx import Document

from sde_value.addbacks import rejection_reason
from sde_value.cli import main
from sde_value.deal import AddBack, load_deal, parse_deal
from sde_value.engine import VERDICTS, value_deal
from sde_value.memo import bottom_line, memo_facts, write_memo

from conftest import ROOT


def raw_hvac():
    with open(ROOT / "deals" / "hvac_service.yaml") as fh:
        return yaml.safe_load(fh)


def memo_text(path):
    doc = Document(path)
    paras = [p.text for p in doc.paragraphs]
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in doc.tables]
    return paras, tables


# --- reasons -----------------------------------------------------------------

def test_reason_for_cost_that_continues():
    ab = AddBack(item="Truck", line="other_deductions", label="discretionary", amount=1,
                 continues=True, note="Used for service calls")
    assert rejection_reason(ab) == ("The cost carries on under a new owner. "
                                    "Used for service calls.")


def test_reason_for_repeated_one_time_item():
    ab = AddBack(item="Repairs", line="repairs", label="one-time", amount=1, years_seen=2)
    assert rejection_reason(ab) == "Called one-time but it appears in 2 of the last 3 years."


def test_reason_for_undocumented_item():
    ab = AddBack(item="Trip", line="advertising", label="discretionary", amount=1,
                 documented=False)
    assert rejection_reason(ab) == "No documents back up a discretionary expense."


def test_reason_for_buyer_override_and_accepted_items():
    salary = AddBack(item="Owner salary", line="officer_comp", label="owner compensation",
                     amount=1, accept=False)
    assert rejection_reason(salary).startswith("Rejected by the buyer")
    trip = AddBack(item="Trip", line="advertising", label="discretionary", amount=1,
                   documented=False, accept=True)
    assert rejection_reason(trip) is None


# --- memo facts --------------------------------------------------------------

def test_rejected_items_add_up_to_the_engine_total(deal_path):
    f = memo_facts(load_deal(deal_path))
    assert sum(r.amount for r in f.rejected) == pytest.approx(f.outputs["rejected_total"])
    assert len(f.rejected) == f.outputs["rejected_count"]
    assert all(r.reason for r in f.rejected)


def test_hvac_rejections(hvac):
    f = memo_facts(hvac)
    assert [r.item for r in f.rejected] == [
        "Spouse on payroll (bookkeeping)", "Owner F-250 lease and fuel",
        "Compressor warranty callbacks", "Trade show trip to Las Vegas"]
    warranty = f.rejected[2]
    assert warranty.classification == "recurring-disguised"
    assert "3 of the last 3 years" in warranty.reason


def test_supported_value_is_capped_by_financing(hvac):
    o = value_deal(hvac)
    assert o["supported_value"] == pytest.approx(o["fin_max_price"])
    assert o["supported_basis"] == "largest financeable price"
    assert o["financing_verdict"] == VERDICTS["neither"]
    line = bottom_line(memo_facts(hvac))
    assert "$1,547,200" in line and "$593,621, 61.6% below the asking price" in line


def test_a_deal_priced_within_reach():
    # The same business offered at 1.0x SDE: the bank can lend on the asking price.
    raw = raw_hvac()
    raw["asking"]["multiple"] = 1.0
    deal = parse_deal(raw)
    o = value_deal(deal)
    assert o["financing_verdict"] == VERDICTS["asking"]
    assert o["supported_basis"] == "repriced value"
    assert o["supported_value"] == pytest.approx(o["repriced_value"])
    assert "The cash flow can finance that" in bottom_line(memo_facts(deal))


def test_overrides_are_listed(tmp_path):
    raw = raw_hvac()
    for ab in raw["addbacks"]:
        if "Las Vegas" in ab["item"]:
            ab["accept"] = True
    f = memo_facts(parse_deal(raw))
    assert f.overridden == ["Trade show trip to Las Vegas"]
    paras, _ = memo_text(write_memo(parse_deal(raw), tmp_path / "memo.docx"))
    assert any("Accepted although the rules would reject them: Trade show trip" in p
               for p in paras)


# --- the document ------------------------------------------------------------

def test_memo_document(tmp_path, hvac):
    path = write_memo(hvac, tmp_path / "hvac-memo.docx")
    paras, tables = memo_text(path)
    assert paras[0] == "Buyer memo: Lone Star Comfort Air (sample)"
    for heading in ("Bottom line", "Asking vs supported value", "Rejected add-backs",
                    "Financing"):
        assert heading in paras
    values, bridge, rejected = tables
    assert values[-1][:3] == ["Supported value", "$593,621", "-61.6%"]
    assert bridge[-1] == ["Normalised EBITDA", "$156,055"]
    assert [row[0] for row in rejected[1:-1]] == [r.item for r in memo_facts(hvac).rejected]
    assert rejected[-1][:2] == ["Total", "$92,700"]
    verdict = next(p for p in paras if p.startswith("Verdict:"))
    assert "neither price can be financed" in verdict and "0.48x" in verdict
    assert any("not an appraisal, tax structuring or legal advice" in p for p in paras)


def test_memo_without_rejections(tmp_path):
    raw = raw_hvac()
    for ab in raw["addbacks"]:
        ab["accept"] = True
    paras, tables = memo_text(write_memo(parse_deal(raw), tmp_path / "memo.docx"))
    assert "None: every add-back the seller claims is accepted." in paras
    assert len(tables) == 2
    assert "Every add-back holds up" in paras[paras.index("Bottom line") + 1]


def test_memo_command(tmp_path, capsys):
    out = tmp_path / "m.docx"
    assert main(["memo", "--deal", str(ROOT / "deals" / "laundromat.yaml"),
                 "--out", str(out)]) == 0
    assert out.exists() and f"Wrote {out}" in capsys.readouterr().out
