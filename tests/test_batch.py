"""The batch command, the summary workbook and the dashboard PDF."""

import re
import shutil

import pytest
from openpyxl import load_workbook

from sde_value.batch import COL, COLUMNS, run_batch
from sde_value.check import check_workbook, recalculate
from sde_value.cli import main
from sde_value.deal import DealError, load_deal
from sde_value.engine import value_deal
from sde_value.export import export_dashboard
from sde_value.workbook import build_workbook

from conftest import DEALS, ROOT

needs_lo = pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")),
                              reason="LibreOffice is required for the recalculation check")

FIRST = 5   # first deal row under the header on the summary sheet


def test_batch_writes_a_workbook_and_memo_per_deal(tmp_path):
    batch = run_batch(ROOT / "deals", tmp_path)
    assert not batch.errors
    assert [r.deal.slug for r in batch.results] == [p.stem for p in DEALS]
    for p in DEALS:
        assert (tmp_path / f"{p.stem}.xlsx").exists()
        assert (tmp_path / f"{p.stem}-memo.docx").exists()
    assert batch.summary == tmp_path / "summary.xlsx"


def test_summary_rows_and_links(tmp_path):
    run_batch(ROOT / "deals", tmp_path)
    ws = load_workbook(tmp_path / "summary.xlsx")["Summary"]
    assert [ws.cell(row=4, column=i).value for i in range(1, len(COLUMNS) + 1)] == \
        [label for _, label, *_ in COLUMNS]
    for r, path in enumerate(DEALS, FIRST):
        o = value_deal(load_deal(path))
        assert ws[f"{COL['file']}{r}"].value == path.name
        assert ws[f"{COL['seller_sde']}{r}"].value == pytest.approx(o["seller_sde"])
        assert ws[f"{COL['financing_verdict']}{r}"].value == o["financing_verdict"]
        # derived columns are formulas, not pasted numbers
        assert ws[f"{COL['supported_value']}{r}"].value.startswith("=MAX(0,MIN(")
        for key in ("workbook", "memo"):
            target = ws[f"{COL[key]}{r}"].hyperlink.target
            assert (tmp_path / target).exists()


@needs_lo
def test_summary_formulas_match_engine_after_recalculation(tmp_path):
    run_batch(ROOT / "deals", tmp_path)
    ws = load_workbook(recalculate(tmp_path / "summary.xlsx", tmp_path / "recalc"),
                       data_only=True)["Summary"]
    gaps = []
    for r, path in enumerate(DEALS, FIRST):
        o = value_deal(load_deal(path))
        for key in ("asking_value", "repriced_value", "supported_value"):
            assert ws[f"{COL[key]}{r}"].value == pytest.approx(o[key], abs=0.005), (path, key)
        got = ws[f"{COL['supported_vs_asking_pct']}{r}"].value
        assert got == pytest.approx(o["supported_vs_asking_pct"], abs=1e-9)
        gaps.append(got)
    avg = ws[f"{COL['supported_vs_asking_pct']}{FIRST + len(DEALS)}"].value
    assert avg == pytest.approx(sum(gaps) / len(gaps), abs=1e-9)
    # supported value sits 55-70% under asking across the samples
    assert avg == pytest.approx(-0.621, abs=0.0005)


def test_a_bad_deal_file_is_skipped_not_fatal(tmp_path, capsys):
    deals = tmp_path / "deals"
    deals.mkdir()
    shutil.copy(ROOT / "deals" / "laundromat.yaml", deals)
    (deals / "broken.yaml").write_text("name: Broken\npnl: {revenue: 100}\n")
    (deals / "README.txt").write_text("not a deal")
    out = tmp_path / "out"
    assert main(["batch", str(deals), "--out", str(out)]) == 1
    captured = capsys.readouterr()
    assert "skipped" in captured.err and "broken.yaml" in captured.err
    assert "Wrote 1 workbook, 1 memo" in captured.out
    assert (out / "laundromat.xlsx").exists() and (out / "summary.xlsx").exists()


def test_an_empty_folder_is_an_error(tmp_path):
    with pytest.raises(DealError, match="no .yaml deal files"):
        run_batch(tmp_path, tmp_path / "out")
    assert main(["batch", str(tmp_path)]) == 2


@needs_lo
def test_batch_check_recalculates_every_workbook(tmp_path, capsys):
    assert main(["batch", str(ROOT / "deals"), "--out", str(tmp_path), "--check"]) == 0
    out = capsys.readouterr().out
    assert out.count(": OK,") == len(DEALS)
    assert "Supported" in out and "$593,621" in out


# --- dashboard ----------------------------------------------------------------

@needs_lo
def test_dashboard_cells_match_engine(tmp_path, hvac):
    path = build_workbook(hvac, tmp_path / "hvac.xlsx")
    wb = load_workbook(path)
    assert wb.active.title == "Dashboard"
    result = check_workbook(path, tabs=["Dashboard"])
    assert result.ok, (result.mismatches, result.missing)
    assert result.compared == 5
    assert result.workbook["o_rejected_count"] == 4
    assert result.workbook["o_supported_basis"] == "largest financeable price"


@needs_lo
def test_dashboard_follows_toggles(tmp_path, hvac):
    """Accept every add-back in Excel: the count, value and verdict move with it."""
    path = build_workbook(hvac, tmp_path / "hvac.xlsx")
    wb = load_workbook(path)
    for i in range(1, len(hvac.addbacks) + 1):
        sheet, ref = next(iter(wb.defined_names[f"i_ab{i}_accept"].destinations))
        wb[sheet][ref.replace("$", "")].value = 1
    wb.save(path)
    result = check_workbook(path, tabs=["Dashboard"])
    assert result.ok, result.mismatches
    assert result.workbook["o_rejected_count"] == 0


@needs_lo
def test_dashboard_pdf_is_one_page(tmp_path, hvac):
    path = build_workbook(hvac, tmp_path / "hvac.xlsx")
    pdf = export_dashboard(path, tmp_path / "dash.pdf")
    data = pdf.read_bytes()
    assert data.startswith(b"%PDF")
    # only the dashboard prints: the other tabs are hidden in the copy sent to LibreOffice
    assert len(re.findall(rb"/Type\s*/Page(?!s)", data)) == 1
    # the export timestamps are blanked out
    assert not re.search(rb"D:\d{8}|<xmp:CreateDate>\d", data)
    assert main(["dashboard", str(path), "--pdf", str(tmp_path / "cli.pdf")]) == 0
